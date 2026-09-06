from __future__ import annotations

import csv
import io
import os
import tempfile
import threading
import uuid
import webbrowser
import zipfile
from pathlib import Path
from typing import Dict, List

from flask import Flask, request, render_template_string, send_file, redirect, url_for
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side

from grader import grade_workbooks, results_to_csv_bytes, results_to_xlsx_bytes

APP_TITLE = "Excel Auto Grader — Final v6 Batch"
PORT = int(os.environ.get("EXCEL_AUTO_GRADER_PORT", "8765"))
HOST = os.environ.get("EXCEL_AUTO_GRADER_HOST", "0.0.0.0")
REPORTS: Dict[str, Dict[str, object]] = {}

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 500 * 1024 * 1024  # 500 MB batch upload limit

INDEX_HTML = """
<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{{ title }}</title>
  <style>
    body { font-family: Arial, sans-serif; margin: 0; background: #f4f6f8; color: #1f2937; }
    .wrap { max-width: 1040px; margin: 32px auto; background: white; padding: 28px; border-radius: 14px; box-shadow: 0 8px 28px rgba(0,0,0,.08); }
    h1 { margin-top: 0; color: #0f172a; }
    .hint { color: #4b5563; line-height: 1.55; }
    .grid { display: grid; grid-template-columns: 1fr 1fr; gap: 18px; margin-top: 24px; }
    .card { border: 1px solid #e5e7eb; border-radius: 10px; padding: 16px; background: #fbfdff; }
    label { display: block; font-weight: bold; margin-bottom: 8px; }
    input[type=file] { width: 100%; box-sizing: border-box; padding: 10px; background: white; border: 1px solid #d1d5db; border-radius: 8px; }
    button { margin-top: 24px; background: #2563eb; color: white; border: none; padding: 12px 20px; border-radius: 10px; font-size: 16px; cursor: pointer; }
    button:hover { background: #1d4ed8; }
    .note { background: #fff7ed; border: 1px solid #fed7aa; padding: 12px; border-radius: 10px; margin-top: 18px; }
    .success { background: #ecfdf5; border: 1px solid #a7f3d0; padding: 12px; border-radius: 10px; margin-top: 18px; }
    .error { background: #fee2e2; border: 1px solid #fca5a5; padding: 12px; border-radius: 10px; margin-top: 18px; }
    .footer { margin-top: 28px; color: #6b7280; font-size: 13px; }
  </style>
</head>
<body>
  <div class="wrap">
    <h1>{{ title }}</h1>
    <p class="hint">Upload one rubric, one model-answer workbook, the original <strong>student starter/template workbook</strong>, and <strong>multiple student workbooks</strong>. Each student file is graded independently. The starter template is used for v6 partial-credit checks such as changed-but-wrong fonts and wrong-target merge attempts.</p>

    {% if error %}<div class="error"><strong>Error:</strong> {{ error }}</div>{% endif %}

    <form method="post" action="/grade-batch" enctype="multipart/form-data">
      <div class="grid">
        <div class="card">
          <label>Rubric Excel file (.xlsx)</label>
          <input type="file" name="rubric" accept=".xlsx" required>
        </div>
        <div class="card">
          <label>Model Answer Excel file (.xlsx)</label>
          <input type="file" name="model" accept=".xlsx" required>
        </div>
        <div class="card">
          <label>Student Completed Excel files (.xlsx) — select multiple</label>
          <input type="file" name="students" accept=".xlsx" multiple required>
        </div>
        <div class="card">
          <label>Original Student Template / Starter Excel file (.xlsx)</label>
          <input type="file" name="template" accept=".xlsx" required>
          <p class="hint" style="font-size:13px">Required for v6 template-aware partial credit (font changes and new/wrong-target merges).</p>
        </div>
      </div>
      <button type="submit">Grade All Student Files</button>
    </form>

    <div class="note">
      <strong>Batch behavior:</strong> files are processed one-by-one. A problem in one student file does not stop the rest of the batch. Failed files are shown in the summary with an error message.
    </div>
    <div class="footer">Codespaces/local server port: {{ port }}</div>
  </div>
</body>
</html>
"""

BATCH_RESULT_HTML = """
<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>{{ title }} - Batch Results</title>
  <style>
    body { font-family: Arial, sans-serif; margin: 0; background: #f4f6f8; color: #1f2937; }
    .wrap { max-width: 1220px; margin: 32px auto; background: white; padding: 28px; border-radius: 14px; box-shadow: 0 8px 28px rgba(0,0,0,.08); }
    h1 { margin-top: 0; color: #0f172a; }
    .stats { display:flex; flex-wrap:wrap; gap:12px; margin: 16px 0 22px; }
    .stat { padding:12px 16px; background:#eff6ff; border:1px solid #bfdbfe; border-radius:10px; }
    .actions a { display:inline-block; margin:0 8px 10px 0; background:#2563eb; color:white; text-decoration:none; padding:10px 14px; border-radius:9px; }
    .actions a.secondary { background:#475569; }
    table { border-collapse:collapse; width:100%; margin-top:20px; font-size:14px; }
    th,td { border:1px solid #e5e7eb; padding:9px; vertical-align:top; }
    th { background:#1e293b; color:white; text-align:left; }
    tr.error td { background:#fff1f2; }
    .good { color:#166534; font-weight:bold; }
    .warn { color:#b45309; font-weight:bold; }
    .bad { color:#b91c1c; font-weight:bold; }
    .small { color:#6b7280; font-size:12px; }
    .download { white-space:nowrap; }
  </style>
</head>
<body>
<div class="wrap">
  <h1>Batch Grading Results</h1>
  <div class="stats">
    <div class="stat"><strong>{{ batch.students_total }}</strong> files selected</div>
    <div class="stat"><strong>{{ batch.students_graded }}</strong> graded successfully</div>
    <div class="stat"><strong>{{ batch.students_failed }}</strong> failed</div>
    <div class="stat">Class average: <strong>{{ batch.average_percentage }}%</strong></div>
  </div>
  <div class="actions">
    <a href="/download/{{ batch_id }}/batch_xlsx">Download Batch Summary XLSX</a>
    <a href="/download/{{ batch_id }}/batch_csv">Download Batch Summary CSV</a>
    <a href="/download/{{ batch_id }}/all_reports_zip">Download All Individual Reports ZIP</a>
    <a class="secondary" href="/">Grade Another Batch</a>
  </div>

  <table>
    <thead><tr>
      <th>#</th><th>Student File</th><th>Score</th><th>Percentage</th><th>Needs Review</th><th>Incorrect</th><th>Partial</th><th>Status</th><th>Detailed Report</th>
    </tr></thead>
    <tbody>
      {% for r in batch.rows %}
      <tr class="{{ 'error' if r.status == 'Error' else '' }}">
        <td>{{ loop.index }}</td>
        <td>{{ r.student_file }}</td>
        <td>{{ r.score_display }}</td>
        <td>{{ r.percentage_display }}</td>
        <td>{{ r.needs_review }}</td>
        <td>{{ r.incorrect }}</td>
        <td>{{ r.partial }}</td>
        <td class="{{ 'bad' if r.status == 'Error' else ('warn' if r.needs_review else 'good') }}">{{ r.status }}{% if r.error %}<br><span class="small">{{ r.error }}</span>{% endif %}</td>
        <td class="download">{% if r.individual_key %}<a href="/download/{{ batch_id }}/individual/{{ r.individual_key }}">Download XLSX</a>{% else %}—{% endif %}</td>
      </tr>
      {% endfor %}
    </tbody>
  </table>
</div>
</body>
</html>
"""


def _safe_filename(name: str) -> str:
    name = Path(name or "student.xlsx").name
    stem = re_sub_invalid(Path(name).stem)
    return (stem or "student") + ".xlsx"


def re_sub_invalid(text: str) -> str:
    import re
    return re.sub(r"[^A-Za-z0-9 _.-]+", "_", text).strip()


def _batch_csv_bytes(rows: List[dict]) -> bytes:
    fields = ["Student_File","Score","Marks_Available","Percentage","Needs_Review","Incorrect","Partial_Credit","Status","Error"]
    sio = io.StringIO()
    w = csv.DictWriter(sio, fieldnames=fields)
    w.writeheader()
    for r in rows:
        w.writerow({
            "Student_File": r["student_file"],
            "Score": r.get("score", ""),
            "Marks_Available": r.get("available", ""),
            "Percentage": r.get("percentage", ""),
            "Needs_Review": r.get("needs_review", ""),
            "Incorrect": r.get("incorrect", ""),
            "Partial_Credit": r.get("partial", ""),
            "Status": r.get("status", ""),
            "Error": r.get("error", ""),
        })
    return sio.getvalue().encode("utf-8-sig")


def _batch_xlsx_bytes(rows: List[dict], batch: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Batch Summary"
    ws["A1"] = "Excel Auto-Grader — Batch Summary"
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws.merge_cells("A1:I1")

    stats = [
        ("Files selected", batch["students_total"]),
        ("Graded successfully", batch["students_graded"]),
        ("Failed", batch["students_failed"]),
        ("Class average", batch["average_percentage"] / 100 if batch["students_graded"] else 0),
    ]
    for idx,(label,value) in enumerate(stats, start=3):
        ws.cell(idx,1,label).font = Font(bold=True)
        ws.cell(idx,2,value)
    ws["B6"].number_format = "0.00%"

    headers = ["Student File","Score","Marks Available","Percentage","Needs Review","Incorrect","Partial Credit","Status","Error"]
    header_row = 8
    for ci,h in enumerate(headers,1):
        c=ws.cell(header_row,ci,h)
        c.font=Font(bold=True,color="FFFFFF")
        c.fill=PatternFill("solid",fgColor="4F81BD")
        c.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)

    for ri,r in enumerate(rows,start=header_row+1):
        vals=[r["student_file"],r.get("score"),r.get("available"),None if r.get("percentage") is None else r["percentage"]/100,r.get("needs_review"),r.get("incorrect"),r.get("partial"),r.get("status"),r.get("error")]
        for ci,v in enumerate(vals,1):
            ws.cell(ri,ci,v)
        if r.get("percentage") is not None:
            ws.cell(ri,4).number_format="0.00%"
        fill = "FFC7CE" if r.get("status")=="Error" else ("FCE4D6" if r.get("needs_review",0) else "C6EFCE")
        ws.cell(ri,8).fill=PatternFill("solid",fgColor=fill)

    widths={"A":38,"B":12,"C":16,"D":14,"E":14,"F":12,"G":14,"H":18,"I":58}
    for col,width in widths.items(): ws.column_dimensions[col].width=width
    thin=Side(style="thin",color="D9E2F3")
    for row in ws.iter_rows(min_row=header_row,max_row=header_row+len(rows),min_col=1,max_col=9):
        for cell in row:
            cell.alignment=Alignment(vertical="top",wrap_text=True)
            cell.border=Border(left=thin,right=thin,top=thin,bottom=thin)
    ws.freeze_panes=f"A{header_row+1}"

    out=io.BytesIO(); wb.save(out); return out.getvalue()


def _reports_zip_bytes(individual_reports: Dict[str, dict]) -> bytes:
    out=io.BytesIO()
    with zipfile.ZipFile(out,"w",zipfile.ZIP_DEFLATED) as z:
        used=set()
        for key,entry in individual_reports.items():
            base=re_sub_invalid(Path(entry["student_file"]).stem) or key
            name=f"{base}_grading_report.xlsx"
            n=2
            while name.casefold() in used:
                name=f"{base}_grading_report_{n}.xlsx"; n+=1
            used.add(name.casefold())
            z.writestr(name,entry["xlsx"])
    return out.getvalue()


@app.route("/", methods=["GET"])
def index():
    return render_template_string(INDEX_HTML, title=APP_TITLE, port=PORT, error=request.args.get("error"))


@app.route("/grade-batch", methods=["POST"])
def grade_batch():
    try:
        rubric = request.files.get("rubric")
        model = request.files.get("model")
        template = request.files.get("template")
        students = [f for f in request.files.getlist("students") if f and f.filename]
        if not rubric or not model or not template or not students:
            return redirect(url_for("index", error="Please upload rubric, model answer, original student template, and at least one student workbook."))
        bad = [f.filename for f in students if Path(f.filename).suffix.lower() != ".xlsx"]
        if bad:
            return redirect(url_for("index", error="Only .xlsx student files are supported in this pilot: " + ", ".join(bad)))

        rows=[]
        individual_reports={}
        with tempfile.TemporaryDirectory() as tmpdir:
            tmp=Path(tmpdir)
            rubric_path=tmp/"rubric.xlsx"; rubric.save(rubric_path)
            model_path=tmp/"model.xlsx"; model.save(model_path)
            template_path=tmp/"student_template.xlsx"; template.save(template_path)

            for idx,student in enumerate(students, start=1):
                original_name=Path(student.filename).name
                student_path=tmp/f"student_{idx}.xlsx"
                student.save(student_path)
                try:
                    results,summary=grade_workbooks(str(rubric_path),str(model_path),str(student_path),student_original_name=original_name,template_path=str(template_path))
                    report_xlsx=results_to_xlsx_bytes(results,summary)
                    key=uuid.uuid4().hex
                    individual_reports[key]={"student_file":original_name,"xlsx":report_xlsx}
                    rows.append({
                        "student_file":original_name,
                        "score":summary["total_awarded"],
                        "available":summary["total_available"],
                        "percentage":summary["percentage"],
                        "needs_review":summary.get("needs_review_count",0),
                        "incorrect":summary.get("incorrect_count",0),
                        "partial":summary.get("partial_count",0),
                        "status":"Needs review" if summary.get("needs_review_count",0) else "Graded",
                        "error":"",
                        "individual_key":key,
                        "score_display":f"{summary['total_awarded']:.2f} / {summary['total_available']:.2f}",
                        "percentage_display":f"{summary['percentage']:.2f}%",
                    })
                except Exception as exc:
                    rows.append({
                        "student_file":original_name,"score":None,"available":None,"percentage":None,
                        "needs_review":0,"incorrect":0,"partial":0,"status":"Error","error":str(exc),
                        "individual_key":"","score_display":"—","percentage_display":"—",
                    })

        ok=[r for r in rows if r["status"]!="Error"]
        batch={
            "students_total":len(rows),
            "students_graded":len(ok),
            "students_failed":len(rows)-len(ok),
            "average_percentage":round(sum(r["percentage"] for r in ok)/len(ok),2) if ok else 0.0,
            "rows":rows,
        }
        batch_id=uuid.uuid4().hex
        REPORTS[batch_id]={
            "batch_xlsx":_batch_xlsx_bytes(rows,batch),
            "batch_csv":_batch_csv_bytes(rows),
            "all_reports_zip":_reports_zip_bytes(individual_reports),
            "individual":individual_reports,
        }
        return render_template_string(BATCH_RESULT_HTML,title=APP_TITLE,batch=batch,batch_id=batch_id)
    except Exception as exc:
        return redirect(url_for("index", error=str(exc)))


@app.route("/download/<batch_id>/<kind>")
def download_batch(batch_id: str, kind: str):
    entry=REPORTS.get(batch_id)
    if not entry or kind not in {"batch_xlsx","batch_csv","all_reports_zip"}:
        return redirect(url_for("index",error="Batch report expired or was not found. Please grade again."))
    data=entry[kind]
    names={
        "batch_xlsx":("excel_batch_grading_summary.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        "batch_csv":("excel_batch_grading_summary.csv","text/csv"),
        "all_reports_zip":("excel_individual_grading_reports.zip","application/zip"),
    }
    filename,mimetype=names[kind]
    return send_file(io.BytesIO(data),mimetype=mimetype,as_attachment=True,download_name=filename)


@app.route("/download/<batch_id>/individual/<key>")
def download_individual(batch_id: str, key: str):
    entry=REPORTS.get(batch_id)
    item=(entry or {}).get("individual",{}).get(key) if entry else None
    if not item:
        return redirect(url_for("index",error="Individual report expired or was not found."))
    base=re_sub_invalid(Path(item["student_file"]).stem) or "student"
    return send_file(io.BytesIO(item["xlsx"]),mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",as_attachment=True,download_name=f"{base}_grading_report.xlsx")


def open_browser_once():
    if not os.environ.get("CODESPACES"):
        webbrowser.open(f"http://127.0.0.1:{PORT}")


if __name__ == "__main__":
    threading.Timer(1.0, open_browser_once).start()
    print(f"{APP_TITLE} is running on port {PORT}")
    print("In Codespaces, open port 8765 from the Ports tab. Press Ctrl+C to stop.")
    app.run(host=HOST, port=PORT, debug=False)
