"""Excel Auto Grader core logic (v6 final batch).

Supports value/formula/style/range/file-name/chart checks using a teacher rubric,
a model-answer workbook, an optional starter/template workbook, and a student workbook.
"""
from __future__ import annotations

import csv
import io
import math
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart
try:
    from openpyxl.chart.bar_chart import BarChart3D
except Exception:
    BarChart3D = ()
try:
    from openpyxl.chart.line_chart import LineChart3D
except Exception:
    LineChart3D = ()
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter, range_boundaries

TRUE_VALUES = {"true", "yes", "y", "1", "correct", "enabled"}
FALSE_VALUES = {"false", "no", "n", "0", "incorrect", "disabled", ""}

CHECK_ALIASES = {
    "cell_value": "value", "value": "value", "text": "value",
    "calculated_value": "calculated_value", "formula": "formula",
    "font_bold": "bold", "bold": "bold", "font_size": "font_size", "size": "font_size",
    "font_name": "font_name", "font": "font_name", "font_color": "font_color",
    "italic": "italic", "underline": "underline", "fill_color": "fill_color", "cell_color": "fill_color",
    "alignment": "alignment", "horizontal_alignment": "alignment", "alignment_horizontal": "alignment",
    "vertical_alignment": "alignment_vertical", "alignment_vertical": "alignment_vertical",
    "wrap_text": "wrap_text", "number_format": "number_format", "cell_style": "cell_style",
    "sheet_exists": "sheet_exists", "range_values": "range_values", "merged_cells": "merged_cells",
    "merge_and_center": "merge_and_center", "column_width": "column_width", "row_height": "row_height",
    "border": "border", "all_borders": "all_borders", "file_name_pattern": "file_name_pattern",
    "chart_line": "chart_line", "line_chart": "chart_line", "chart_bar": "chart_bar", "bar_chart": "chart_bar",
}

RANGE_STYLE_CHECKS = {
    "bold", "italic", "underline", "font_size", "font_name", "font_color", "fill_color",
    "alignment", "alignment_vertical", "wrap_text", "number_format", "cell_style", "border",
}

REPORT_HEADERS = [
    "Task_ID", "Task_Description", "Sheet", "Target", "Check_Type", "Expected",
    "Student", "First_Pass_Result", "Verification_Result", "Result",
    "Marks_Awarded", "Marks_Available", "Feedback"
]


@dataclass
class TaskResult:
    task_id: str
    task_description: str
    sheet: str
    target: str
    check_type: str
    expected: str
    student: str
    first_pass_result: str
    verification_result: str
    result: str
    marks_awarded: float
    marks_available: float
    feedback: str

    def to_row(self) -> Dict[str, Any]:
        return {
            "Task_ID": self.task_id, "Task_Description": self.task_description, "Sheet": self.sheet,
            "Target": self.target, "Check_Type": self.check_type, "Expected": self.expected,
            "Student": self.student, "First_Pass_Result": self.first_pass_result,
            "Verification_Result": self.verification_result, "Result": self.result,
            "Marks_Awarded": self.marks_awarded, "Marks_Available": self.marks_available,
            "Feedback": self.feedback,
        }


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"[^a-z0-9]+", "_", str(value).strip().lower()).strip("_")


def clean_string(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_uploaded_filename(name: str, mode: str = "strip_upload_copy_suffix") -> str:
    """Normalize browser/upload duplicate suffixes such as file(2).xlsx or file (2).xlsx."""
    base = os.path.basename(name or "")
    mode = clean_string(mode).lower() or "strip_upload_copy_suffix"
    if mode in {"strip_upload_copy_suffix", "strip_copy_suffix", "normalized"}:
        stem, ext = os.path.splitext(base)
        stem = re.sub(r"\s*\(\d+\)$", "", stem).rstrip()
        return stem + ext
    return base


def is_enabled(value: Any) -> bool:
    if value is None:
        return False
    if isinstance(value, bool):
        return value
    s = clean_string(value).lower()
    if s in TRUE_VALUES:
        return True
    if s in FALSE_VALUES:
        return False
    return True


def as_float(value: Any, default: float = 0.0) -> float:
    if value in (None, ""):
        return default
    try:
        return float(value)
    except Exception:
        return default


def normalize_formula(formula: Any, mode: str = "ignore_spaces") -> str:
    text = clean_string(formula)
    if not text:
        return ""
    if not text.startswith("="):
        text = "=" + text
    if mode in {"ignore_spaces", "ignore_space", "normalized"}:
        text = re.sub(r"\s+", "", text)
    if mode in {"ignore_spaces", "ignore_case", "normalized"}:
        text = text.upper()
    return text


def formula_structure_signature(formula: Any, mode: str = "ignore_spaces") -> str:
    """Abstract cell/range references while preserving functions/operators for partial formula matching."""
    text = normalize_formula(formula, mode)
    if not text:
        return ""
    # Normalize sheet-qualified and ordinary A1 references/ranges to REF tokens.
    text = re.sub(r"(?:'[^']+'|[A-Z0-9_ .]+)!\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?", "REF", text)
    text = re.sub(r"\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?", "REF", text)
    # Normalize numeric constants so a minor constant error can still be recognized structurally.
    text = re.sub(r"(?<![A-Z])\d+(?:\.\d+)?", "NUM", text)
    return text


def formula_partial_factor(expected_formula: Any, student_formula: Any, model_calc: Any = None, student_calc: Any = None, tolerance: float = 0.0, mode: str = "ignore_spaces") -> float:
    """1.0 exact/equivalent, 0.5 meaningful partial match, otherwise 0.0."""
    exp = normalize_formula(expected_formula, mode)
    stu = normalize_formula(student_formula, mode)
    if not exp or not stu:
        return 0.0
    if exp == stu:
        return 1.0
    # Correct calculated result is accepted as a meaningful partial formula attempt.
    if model_calc is not None and student_calc is not None and compare_values(model_calc, student_calc, "exact", tolerance):
        return 0.5
    # Same abstract formula structure (same function/operator pattern with different refs/constants).
    if formula_structure_signature(exp, mode) == formula_structure_signature(stu, mode):
        return 0.5
    # Same main Excel function is also a meaningful partial attempt for copied MAX/AVERAGE/etc.
    ef = re.match(r"=([A-Z][A-Z0-9_.]*)\(", exp)
    sf = re.match(r"=([A-Z][A-Z0-9_.]*)\(", stu)
    if ef and sf and ef.group(1) == sf.group(1):
        return 0.5
    # For arithmetic formulas, preserve the operator sequence as a conservative structural check.
    eops = re.findall(r"[+\-*/^]", exp)
    sops = re.findall(r"[+\-*/^]", stu)
    if eops and eops == sops:
        return 0.5
    return 0.0


def normalize_bool_text(value: Any) -> str:
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    s = clean_string(value).lower()
    if s in TRUE_VALUES:
        return "TRUE"
    if s in FALSE_VALUES:
        return "FALSE"
    return clean_string(value)


def color_signature(color: Any) -> str:
    if color is None:
        return ""
    ctype = getattr(color, "type", None)
    tint = getattr(color, "tint", 0.0) or 0.0
    if ctype == "rgb" and getattr(color, "rgb", None):
        return f"RGB:{str(color.rgb).upper()}:TINT:{round(float(tint), 6)}"
    if ctype == "theme" and getattr(color, "theme", None) is not None:
        return f"THEME:{color.theme}:TINT:{round(float(tint), 6)}"
    if ctype == "indexed" and getattr(color, "indexed", None) is not None:
        return f"INDEXED:{color.indexed}"
    if ctype == "auto" and getattr(color, "auto", None) is not None:
        return f"AUTO:{bool(color.auto)}"
    rgb = getattr(color, "rgb", None)
    if rgb:
        return f"RGB:{str(rgb).upper()}"
    return clean_string(color)


def display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float) and math.isclose(value, round(value), rel_tol=0, abs_tol=1e-12):
        return str(int(round(value)))
    return str(value)


def compare_values(expected: Any, student: Any, comparison_mode: str = "exact", tolerance: float = 0.0) -> bool:
    mode = clean_string(comparison_mode).lower() or "exact"
    try:
        e_num, s_num = float(expected), float(student)
        if not (math.isnan(e_num) or math.isnan(s_num)):
            return abs(e_num - s_num) <= tolerance
    except Exception:
        pass
    e, s = "" if expected is None else str(expected), "" if student is None else str(student)
    if mode in {"case_insensitive", "ignore_case"}:
        return e.strip().lower() == s.strip().lower()
    if mode in {"contains", "student_contains_expected"}:
        return e.strip().lower() in s.strip().lower()
    if mode == "model_contains_student":
        return s.strip().lower() in e.strip().lower()
    if mode in {"ignore_spaces", "normalized"}:
        return re.sub(r"\s+", "", e).lower() == re.sub(r"\s+", "", s).lower()
    if mode == "regex":
        try:
            return re.fullmatch(e, s, flags=re.IGNORECASE) is not None
        except re.error:
            return False
    return e.strip() == s.strip()


def find_rubric_sheet(wb) -> Any:
    for name in ["Rubric", "rubric", "Tasks", "tasks"]:
        if name in wb.sheetnames:
            return wb[name]
    return wb[wb.sheetnames[0]]


def read_rubric(rubric_path: str) -> List[Dict[str, Any]]:
    wb = load_workbook(rubric_path, data_only=True)
    ws = find_rubric_sheet(wb)
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header_row_idx = None
    header_map: Dict[int, str] = {}
    for idx, row in enumerate(rows[:20]):
        normalized = [normalize_header(x) for x in row]
        if "task_id" in normalized or "task_description" in normalized or "marks" in normalized:
            header_row_idx = idx
            header_map = {i: h for i, h in enumerate(normalized) if h}
            break
    if header_row_idx is None:
        raise ValueError("Could not find rubric header row.")
    tasks = []
    for row in rows[header_row_idx + 1:]:
        record = {header_map.get(i, f"col_{i}"): value for i, value in enumerate(row) if header_map.get(i)}
        if not any(v not in (None, "") for v in record.values()):
            continue
        if "enabled" in record and not is_enabled(record.get("enabled")):
            continue
        if not clean_string(record.get("task_id")) and not clean_string(record.get("task_description")):
            continue
        tasks.append(record)
    return tasks


def resolve_sheet_name(task: Dict[str, Any], wb) -> str:
    lookup = clean_string(task.get("sheet_lookup")).lower() or "sheet_name"
    name = clean_string(task.get("sheet_name"))
    number_raw = task.get("sheet_number")
    if lookup in {"sheet_number", "number", "index", "sheet_index"} or (not name and number_raw not in (None, "")):
        idx = int(as_float(number_raw, 1)) - 1
        if idx < 0 or idx >= len(wb.sheetnames):
            raise KeyError(f"Sheet number {number_raw} not found")
        return wb.sheetnames[idx]
    if name:
        if name in wb.sheetnames:
            return name
        lower_map = {s.lower(): s for s in wb.sheetnames}
        if name.lower() in lower_map:
            return lower_map[name.lower()]
        raise KeyError(f"Sheet '{name}' not found")
    return wb.sheetnames[0]


def resolve_target(task: Dict[str, Any]) -> str:
    address = clean_string(task.get("target_address"))
    if address:
        return address
    row, col = task.get("row_number"), task.get("column_number")
    if row not in (None, "") and col not in (None, ""):
        return f"{get_column_letter(int(as_float(col, 1)))}{int(as_float(row, 1))}"
    return ""


def infer_check_type(task: Dict[str, Any]) -> str:
    raw = clean_string(task.get("check_type")).lower()
    if raw:
        return CHECK_ALIASES.get(raw, raw)
    desc = clean_string(task.get("task_description")).lower()
    if "line chart" in desc: return "chart_line"
    if "bar chart" in desc: return "chart_bar"
    if "formula" in desc or "function" in desc: return "formula"
    if "wrap" in desc: return "wrap_text"
    if "bold" in desc: return "bold"
    if "font size" in desc: return "font_size"
    if "font color" in desc or "font colour" in desc: return "font_color"
    if "align" in desc or "center" in desc or "centre" in desc: return "alignment"
    if "fill" in desc or "color" in desc or "colour" in desc: return "fill_color"
    return "value"


def iter_cells(ws, address: str):
    min_col, min_row, max_col, max_row = range_boundaries(address)
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            yield cell


def cell_property(cell, check_type: str):
    if check_type == "bold": return bool(cell.font and cell.font.bold)
    if check_type == "italic": return bool(cell.font and cell.font.italic)
    if check_type == "underline": return bool(cell.font and cell.font.underline)
    if check_type == "font_size": return cell.font.sz if cell.font else None
    if check_type == "font_name": return cell.font.name if cell.font else None
    if check_type == "font_color": return color_signature(cell.font.color if cell.font else None)
    if check_type == "fill_color":
        fill = cell.fill
        if fill is None or fill.fill_type in (None, "none"):
            return ""
        return color_signature(fill.fgColor)
    if check_type == "alignment": return cell.alignment.horizontal if cell.alignment else None
    if check_type == "alignment_vertical": return cell.alignment.vertical if cell.alignment else None
    if check_type == "wrap_text": return bool(cell.alignment and cell.alignment.wrap_text)
    if check_type == "number_format": return cell.number_format
    if check_type == "cell_style": return cell.style
    if check_type == "border":
        b = cell.border
        return bool(b and any(side and side.style for side in [b.left, b.right, b.top, b.bottom]))
    return cell.value


def border_signature(cell) -> Tuple[str, str, str, str]:
    b = cell.border
    if not b:
        return ("", "", "", "")
    return tuple((getattr(side, "style", None) or "") for side in [b.left, b.right, b.top, b.bottom])


def get_property_or_range(wb_formula, sheet: str, address: str, check_type: str):
    ws = wb_formula[sheet]
    if ":" not in address:
        return cell_property(ws[address], check_type)
    return [cell_property(c, check_type) for c in iter_cells(ws, address)]


def get_cell_value(wb_formula, wb_values, sheet: str, address: str, check_type: str):
    ws_f, ws_v = wb_formula[sheet], wb_values[sheet]
    if check_type == "formula": return ws_f[address].value
    if check_type in {"calculated_value", "value"}: return ws_v[address].value
    if check_type in RANGE_STYLE_CHECKS: return get_property_or_range(wb_formula, sheet, address, check_type)
    if check_type == "column_width": return ws_f.column_dimensions[address.upper()].width
    if check_type == "row_height": return ws_f.row_dimensions[int(address)].height
    return ws_v[address].value


def get_range_values(wb_values, sheet: str, address: str) -> List[List[Any]]:
    ws = wb_values[sheet]
    min_col, min_row, max_col, max_row = range_boundaries(address)
    return [[cell.value for cell in row] for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col)]


def get_merged_ranges(wb_formula, sheet: str) -> List[str]:
    return [str(rng) for rng in wb_formula[sheet].merged_cells.ranges]


def merge_and_center_signature(wb_formula, sheet: str, address: str) -> Dict[str, Any]:
    ws = wb_formula[sheet]
    target_norm = address.replace("$", "").upper()
    merged = any(str(r).replace("$", "").upper() == target_norm for r in ws.merged_cells.ranges)
    min_col, min_row, _, _ = range_boundaries(address)
    top_left = ws.cell(min_row, min_col)
    centered = (top_left.alignment.horizontal or "").lower() in {"center", "centercontinuous"}
    return {"merged": merged, "centered": centered}


def all_borders_signature(wb_formula, sheet: str, address: str) -> List[Tuple[str, str, str, str]]:
    return [border_signature(c) for c in iter_cells(wb_formula[sheet], address)]


def normalize_chart_ref(ref: Optional[str]) -> str:
    if not ref:
        return ""
    return re.sub(r"[\s$']", "", str(ref)).upper()


def nested_ref_formula(obj: Any) -> str:
    if obj is None:
        return ""
    for attr in ("numRef", "strRef", "multiLvlStrRef"):
        nested = getattr(obj, attr, None)
        if nested is not None:
            f = getattr(nested, "f", None)
            if f:
                return normalize_chart_ref(f)
    return ""


def series_signature(series: Any) -> Tuple[str, str]:
    cat = nested_ref_formula(getattr(series, "cat", None)) or nested_ref_formula(getattr(series, "xVal", None))
    val = nested_ref_formula(getattr(series, "val", None)) or nested_ref_formula(getattr(series, "yVal", None))
    return (cat, val)


def chart_dimension(chart: Any) -> str:
    name = chart.__class__.__name__.lower()
    return "3d" if "3d" in name else "2d"


def chart_kind(chart: Any) -> str:
    if isinstance(chart, LineChart): return "line"
    if isinstance(chart, BarChart):
        direction = clean_string(getattr(chart, "barDir", None) or getattr(chart, "type", None)).lower()
        return "bar" if direction == "bar" else "column"
    name = chart.__class__.__name__.lower()
    if "line" in name: return "line"
    if "bar" in name: return "bar"
    return name


def chart_signatures(wb_formula, sheet: str, wanted: str) -> List[Dict[str, Any]]:
    out = []
    for chart in getattr(wb_formula[sheet], "_charts", []):
        kind = chart_kind(chart)
        if wanted == "chart_line" and kind != "line":
            continue
        if wanted == "chart_bar" and kind != "bar":
            continue
        series = sorted(series_signature(s) for s in getattr(chart, "ser", []))
        out.append({"kind": kind, "dimension": chart_dimension(chart), "series": series})
    return out


def get_expected_value(task, model_formula, model_values, sheet, target, check_type, student_original_name=""):
    source = clean_string(task.get("expected_source")).lower() or "model_answer"
    if source in {"manual", "manual_value", "typed_value", "expected_value"}:
        return task.get("expected_value")
    if check_type == "file_name_pattern":
        return task.get("expected_value") or r"^[^_]+_[^_]+\.xlsx$"
    if check_type == "sheet_exists": return True
    if check_type == "range_values": return get_range_values(model_values, sheet, target)
    if check_type == "merged_cells": return get_merged_ranges(model_formula, sheet)
    if check_type == "merge_and_center": return merge_and_center_signature(model_formula, sheet, target)
    if check_type == "all_borders": return all_borders_signature(model_formula, sheet, target)
    if check_type in {"chart_line", "chart_bar"}: return chart_signatures(model_formula, sheet, check_type)
    return get_cell_value(model_formula, model_values, sheet, target, check_type)


def get_student_value(task, student_formula, student_values, sheet, target, check_type, student_original_name=""):
    if check_type == "file_name_pattern":
        mode = clean_string(task.get("filename_normalization")) or "strip_upload_copy_suffix"
        return normalize_uploaded_filename(student_original_name or "", mode)
    if check_type == "sheet_exists": return sheet in student_formula.sheetnames
    if check_type == "range_values": return get_range_values(student_values, sheet, target)
    if check_type == "merged_cells": return get_merged_ranges(student_formula, sheet)
    if check_type == "merge_and_center": return merge_and_center_signature(student_formula, sheet, target)
    if check_type == "all_borders": return all_borders_signature(student_formula, sheet, target)
    if check_type in {"chart_line", "chart_bar"}: return chart_signatures(student_formula, sheet, check_type)
    return get_cell_value(student_formula, student_values, sheet, target, check_type)


def chart_equal(expected: List[Dict[str, Any]], student: List[Dict[str, Any]]) -> bool:
    if not expected or not student:
        return False
    # A student only needs one chart that matches one model-answer chart of the requested type.
    for e in expected:
        for s in student:
            if e.get("kind") == s.get("kind") and e.get("series") == s.get("series"):
                return True
    return False


def chart_match_score(expected: List[Dict[str, Any]], student: List[Dict[str, Any]]) -> float:
    """Return 1.0 exact chart match, 0.5 only-2D/3D mismatch, otherwise 0.0."""
    if not expected or not student:
        return 0.0
    half = False
    for e in expected:
        for s in student:
            if e.get("kind") != s.get("kind"):
                continue
            if e.get("series") != s.get("series"):
                continue
            if e.get("dimension") == s.get("dimension"):
                return 1.0
            half = True
    return 0.5 if half else 0.0


def values_equal(expected: Any, student: Any, check_type: str, task: Dict[str, Any]) -> bool:
    comparison_mode = clean_string(task.get("comparison_mode")).lower() or "exact"
    formula_mode = clean_string(task.get("formula_mode")).lower() or "ignore_spaces"
    tolerance = as_float(task.get("tolerance"), 0.0)
    if check_type == "formula": return normalize_formula(expected, formula_mode) == normalize_formula(student, formula_mode)
    if check_type in {"bold", "italic", "underline", "sheet_exists", "border", "wrap_text"} and not isinstance(expected, list):
        return normalize_bool_text(expected) == normalize_bool_text(student)
    if check_type in {"font_size", "column_width", "row_height", "calculated_value", "value"}:
        return compare_values(expected, student, comparison_mode, tolerance)
    if check_type == "file_name_pattern":
        pattern = clean_string(expected) or r"^[^_]+_[^_]+\.xlsx$"
        try:
            return re.fullmatch(pattern, clean_string(student), flags=re.IGNORECASE) is not None
        except re.error:
            return False
    if check_type in {"chart_line", "chart_bar"}: return chart_equal(expected, student)
    if check_type == "merge_and_center": return bool(student.get("merged")) and bool(student.get("centered")) and student == expected
    if check_type in {"all_borders", "range_values"} or isinstance(expected, list):
        if len(expected) != len(student): return False
        for e, s in zip(expected, student):
            if isinstance(e, (tuple, dict, list)):
                if e != s: return False
            elif check_type in {"bold", "italic", "underline", "wrap_text"}:
                if normalize_bool_text(e) != normalize_bool_text(s): return False
            elif check_type == "font_size":
                if not compare_values(e, s, comparison_mode, tolerance): return False
            else:
                if not compare_values(e, s, comparison_mode, tolerance): return False
        return True
    if check_type == "merged_cells": return set(map(str, expected)) == set(map(str, student))
    return compare_values(expected, student, comparison_mode, tolerance)


def stringify_for_report(value: Any) -> str:
    if isinstance(value, (list, dict, tuple)):
        text = str(value)
        return text if len(text) <= 600 else text[:597] + "..."
    return display_value(value)


def grade_workbooks(rubric_path: str, model_path: str, student_path: str, student_original_name: Optional[str] = None, template_path: Optional[str] = None) -> Tuple[List[TaskResult], Dict[str, float]]:
    tasks = read_rubric(rubric_path)
    model_formula = load_workbook(model_path, data_only=False)
    model_values = load_workbook(model_path, data_only=True)
    student_formula = load_workbook(student_path, data_only=False)
    student_values = load_workbook(student_path, data_only=True)
    template_formula = load_workbook(template_path, data_only=False) if template_path else None
    template_values = load_workbook(template_path, data_only=True) if template_path else None
    original_name = student_original_name or Path(student_path).name
    results: List[TaskResult] = []

    for i, task in enumerate(tasks, start=1):
        task_id = clean_string(task.get("task_id")) or f"T{i:03d}"
        desc = clean_string(task.get("task_description")) or task_id
        marks = as_float(task.get("marks"), 0.0)
        check_type = infer_check_type(task)
        target = resolve_target(task)
        try:
            # Filename checks are workbook-level and do not need a sheet, but still resolve one for reporting.
            sheet = resolve_sheet_name(task, model_formula) if model_formula.sheetnames else ""
        except Exception as exc:
            results.append(TaskResult(task_id, desc, clean_string(task.get("sheet_name")), target, check_type, "", "", "Model answer issue", "Model answer issue", "Model answer issue", 0.0, marks, str(exc)))
            continue
        try:
            if check_type not in {"file_name_pattern", "sheet_exists"} and sheet not in student_formula.sheetnames:
                raise KeyError(f"Student workbook does not contain sheet '{sheet}'")
            expected = get_expected_value(task, model_formula, model_values, sheet, target, check_type, original_name)
            student = get_student_value(task, student_formula, student_values, sheet, target, check_type, original_name)
            if check_type == "formula" and not (isinstance(expected, str) and expected.startswith("=")):
                results.append(TaskResult(task_id, desc, sheet, target, check_type, stringify_for_report(expected), stringify_for_report(student), "Model answer issue", "Model answer issue", "Model answer issue", 0.0, marks, "Rubric asks for a formula, but the model-answer cell does not contain a formula."))
                continue
            if check_type in {"chart_line", "chart_bar"} and not expected:
                results.append(TaskResult(task_id, desc, sheet, target, check_type, "", stringify_for_report(student), "Model answer issue", "Model answer issue", "Model answer issue", 0.0, marks, "The model-answer workbook does not contain the required chart type."))
                continue
            grade_mode = clean_string(task.get("grade_mode")).lower() or "all_or_nothing"

            def score_current(mf, mv, sf, sv, tf=None, tv=None, exp_value=None, stu_value=None):
                exp_value = expected if exp_value is None else exp_value
                stu_value = student if stu_value is None else stu_value
                if check_type in {"chart_line", "chart_bar"} and grade_mode == "chart_dimension_half_credit":
                    return chart_match_score(exp_value, stu_value)
                if grade_mode == "font_changed_from_template_partial" and check_type == "font_name":
                    if values_equal(exp_value, stu_value, check_type, task):
                        return 1.0
                    if tf is None:
                        return 0.0
                    template_font = get_cell_value(tf, tv or tf, sheet, target, check_type)
                    return 0.5 if clean_string(stu_value).casefold() != clean_string(template_font).casefold() else 0.0
                if grade_mode == "formula_structural_partial" and check_type == "formula":
                    model_calc = mv[sheet][target].value if target else None
                    student_calc = sv[sheet][target].value if target else None
                    return formula_partial_factor(exp_value, stu_value, model_calc, student_calc, as_float(task.get("tolerance"), 0.0), clean_string(task.get("formula_mode")) or "ignore_spaces")
                if grade_mode == "merge_wrong_target_partial" and check_type == "merge_and_center":
                    if values_equal(exp_value, stu_value, check_type, task):
                        return 1.0
                    if tf is None:
                        return 0.0
                    template_merges = set(get_merged_ranges(tf, sheet))
                    student_merges = set(get_merged_ranges(sf, sheet))
                    new_merges = student_merges - template_merges
                    target_norm = target.replace("$", "").upper()
                    wrong_target_new = [m for m in new_merges if m.replace("$", "").upper() != target_norm]
                    return 0.5 if wrong_target_new else 0.0
                return 1.0 if values_equal(exp_value, stu_value, check_type, task) else 0.0

            # First pass
            score_factor = score_current(model_formula, model_values, student_formula, student_values, template_formula, template_values)
            first_pass = "Correct" if score_factor == 1.0 else ("Partial credit" if score_factor == 0.5 else "Incorrect")

            # Verification pass: re-open the source files and re-read the exact task evidence.
            verification_mode = clean_string(task.get("verification_mode")).lower() or "two_pass"
            verification = first_pass
            if verification_mode in {"two_pass", "verify_incorrect", "always"} and first_pass != "Correct":
                mf2 = load_workbook(model_path, data_only=False)
                mv2 = load_workbook(model_path, data_only=True)
                sf2 = load_workbook(student_path, data_only=False)
                sv2 = load_workbook(student_path, data_only=True)
                expected2 = get_expected_value(task, mf2, mv2, sheet, target, check_type, original_name)
                student2 = get_student_value(task, sf2, sv2, sheet, target, check_type, original_name)
                tf2 = load_workbook(template_path, data_only=False) if template_path else None
                tv2 = load_workbook(template_path, data_only=True) if template_path else None
                score_factor2 = score_current(mf2, mv2, sf2, sv2, tf2, tv2, expected2, student2)
                verification = "Correct" if score_factor2 == 1.0 else ("Partial credit" if score_factor2 == 0.5 else "Incorrect")
                # If the two passes disagree, do not silently assign zero.
                if verification != first_pass:
                    results.append(TaskResult(task_id, desc, sheet, target, check_type, stringify_for_report(expected2), stringify_for_report(student2), first_pass, verification, "Needs review", 0.0, marks, "Two grading passes disagreed. Review this task before finalizing the grade."))
                    continue
                expected, student, score_factor = expected2, student2, score_factor2

            # Formula safeguard: if formula check fails but cached calculated values match, flag for review rather than a silent zero.
            if check_type == "formula" and score_factor == 0.0 and target and grade_mode != "formula_structural_partial":
                try:
                    model_calc = model_values[sheet][target].value
                    student_calc = student_values[sheet][target].value
                    if model_calc is not None and student_calc is not None and compare_values(model_calc, student_calc, "exact", as_float(task.get("tolerance"), 0.0)):
                        results.append(TaskResult(task_id, desc, sheet, target, check_type, stringify_for_report(expected), stringify_for_report(student), first_pass, verification, "Needs review", 0.0, marks, f"Formula text did not match, but calculated values match ({display_value(student_calc)}). Manual review required."))
                        continue
                except Exception:
                    pass

            awarded = round(marks * score_factor, 4)
            final_result = "Correct" if score_factor == 1.0 else ("Partial credit" if score_factor == 0.5 else "Incorrect")
            if final_result == "Correct":
                feedback = ""
            elif final_result == "Partial credit":
                if grade_mode == "chart_dimension_half_credit":
                    feedback = "Correct chart family and source ranges, but 2D/3D dimension differs from the model answer; 50% credit applied."
                elif grade_mode == "font_changed_from_template_partial":
                    feedback = "The student changed the template-default font, but did not use the required/model font; 50% credit applied."
                elif grade_mode == "formula_structural_partial":
                    feedback = "The copied formula meaningfully matches the required formula structure or calculated result, but is not fully equivalent; 50% credit applied."
                elif grade_mode == "merge_wrong_target_partial":
                    feedback = "A new merge was created on other non-default cells rather than correctly completing the required target merge; 50% credit applied."
                else:
                    feedback = clean_string(task.get("comment_feedback")) or "Partial credit applied according to the rubric."
            else:
                feedback = clean_string(task.get("comment_feedback")) or "Expected value does not match student value."
            results.append(TaskResult(task_id, desc, sheet, target, check_type, stringify_for_report(expected), stringify_for_report(student), first_pass, verification, final_result, awarded, marks, feedback))
        except Exception as exc:
            results.append(TaskResult(task_id, desc, sheet, target, check_type, "", "", "Error", "Error", "Error", 0.0, marks, str(exc)))

    total_awarded = sum(r.marks_awarded for r in results)
    total_available = sum(r.marks_available for r in results)
    summary = {
        "total_awarded": round(total_awarded, 4), "total_available": round(total_available, 4),
        "percentage": round((total_awarded / total_available * 100) if total_available else 0.0, 2),
        "tasks_count": len(results),
        "needs_review_count": sum(1 for r in results if r.result == "Needs review"),
        "incorrect_count": sum(1 for r in results if r.result == "Incorrect"),
        "partial_count": sum(1 for r in results if r.result == "Partial credit"),
    }
    return results, summary


def results_to_csv_bytes(results: List[TaskResult], summary: Dict[str, float]) -> bytes:
    buffer = io.StringIO()
    writer = csv.DictWriter(buffer, fieldnames=REPORT_HEADERS)
    writer.writeheader()
    for r in results: writer.writerow(r.to_row())
    writer.writerow({})
    writer.writerow({"Task_ID": "TOTAL", "Marks_Awarded": summary["total_awarded"], "Marks_Available": summary["total_available"], "Result": f"{summary['percentage']}%"})
    return buffer.getvalue().encode("utf-8-sig")


def results_to_xlsx_bytes(results: List[TaskResult], summary: Dict[str, float]) -> bytes:
    wb = Workbook(); ws = wb.active; ws.title = "Grading Report"
    ws["A1"] = "Excel Auto-Grading Report"; ws["A1"].font = Font(bold=True, size=16)
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "Total Score", summary["total_awarded"], "/", summary["total_available"]
    ws["A4"], ws["B4"] = "Percentage", summary["percentage"] / 100; ws["B4"].number_format = "0.00%"
    start_row = 6
    for col_idx, header in enumerate(REPORT_HEADERS, start=1):
        c = ws.cell(start_row, col_idx, header); c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="4F81BD"); c.alignment = Alignment(horizontal="center")
    for row_idx, r in enumerate(results, start=start_row + 1):
        row = r.to_row()
        for col_idx, header in enumerate(REPORT_HEADERS, start=1): ws.cell(row_idx, col_idx, row[header])
        rc = ws.cell(row_idx, 10)
        if r.result == "Correct": color = "C6EFCE"
        elif r.result == "Partial credit": color = "FFEB9C"
        elif r.result == "Needs review": color = "FCE4D6"
        else: color = "FFC7CE"
        rc.fill = PatternFill("solid", fgColor=color)
    widths = {"A":12,"B":42,"C":16,"D":14,"E":20,"F":38,"G":38,"H":18,"I":18,"J":18,"K":16,"L":16,"M":48}
    for col, width in widths.items(): ws.column_dimensions[col].width = width
    thin = Side(style="thin", color="D9E2F3")
    for row in ws.iter_rows(min_row=start_row, max_row=start_row + len(results), min_col=1, max_col=len(REPORT_HEADERS)):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True); cell.border = Border(left=thin,right=thin,top=thin,bottom=thin)
    ws.freeze_panes = f"A{start_row + 1}"
    output = io.BytesIO(); wb.save(output); return output.getvalue()
